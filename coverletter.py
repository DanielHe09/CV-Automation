#!/usr/bin/env python3
"""Generate a tailored cover letter PDF from the base template."""

import argparse
import datetime
import re
import subprocess
import sys
import tempfile
from pathlib import Path

import docx

BASE_DIR = Path(__file__).resolve().parent
TEMPLATE = BASE_DIR / "Cover Letter Template.docx"
OUTPUT_DIR = Path.home() / "Downloads"
DEFAULT_ROLE = "Software Engineering Internship"
DEFAULT_ADJECTIVE = "innovative"
SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"


def fill_placeholders(document, values):
    for paragraph in document.paragraphs:
        for run in paragraph.runs:
            for key, value in values.items():
                if key in run.text:
                    run.text = run.text.replace(key, value)


def docx_to_pdf(docx_path, pdf_path):
    out_dir = pdf_path.parent
    result = subprocess.run(
        [
            SOFFICE,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(out_dir),
            str(docx_path),
        ],
        capture_output=True,
        text=True,
    )
    if result.returncode != 0:
        sys.exit(f"PDF export failed:\n{result.stderr.strip()}")

    # LibreOffice names the output after the input file; rename to our target name.
    produced = out_dir / (docx_path.stem + ".pdf")
    produced.replace(pdf_path)


def safe_filename(text):
    return re.sub(r"[^\w\s-]", "", text).strip()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("company", help='Company name, e.g. "Shopify"')
    parser.add_argument(
        "role",
        nargs="?",
        default=DEFAULT_ROLE,
        help=f'Role title (default: "{DEFAULT_ROLE}")',
    )
    parser.add_argument(
        "adjective",
        nargs="?",
        default=DEFAULT_ADJECTIVE,
        help=f'Adjective describing the company environment (default: "{DEFAULT_ADJECTIVE}")',
    )
    parser.add_argument(
        "--role",
        dest="role_flag",
        help="Role title, as a flag so you can skip the positional arguments",
    )
    parser.add_argument(
        "--adj",
        dest="adj_flag",
        help="Adjective, as a flag so you can set it without typing the role",
    )
    parser.add_argument(
        "--date",
        help="Date on the letter (default: today, e.g. August 10, 2026)",
    )
    args = parser.parse_args()

    role = args.role_flag or args.role
    adjective = args.adj_flag or args.adjective

    letter_date = args.date or datetime.date.today().strftime("%B %-d, %Y")

    document = docx.Document(TEMPLATE)
    fill_placeholders(
        document,
        {
            "[COMPANY]": args.company,
            "[ROLE]": role,
            "[ADJECTIVE]": adjective,
            "[DATE]": letter_date,
        },
    )

    OUTPUT_DIR.mkdir(exist_ok=True)
    pdf_path = OUTPUT_DIR / f"Daniel He {safe_filename(args.company)} Cover Letter.pdf"

    with tempfile.TemporaryDirectory() as tmp:
        temp_docx = Path(tmp) / "letter.docx"
        document.save(temp_docx)
        docx_to_pdf(temp_docx, pdf_path)

    print(f"Created {pdf_path}")


if __name__ == "__main__":
    main()
